"""Discover and load LLM evaluation artifacts without double-counting runs."""

from __future__ import annotations

import ast
import csv
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from interactive.evaluator import evaluation_coverage, flatten_evaluation_records
from interactive.multilingual import TRACK_UTTERANCES, _norm_key

CLASS_MODEL = "MODEL EVALUATION"
CLASS_MULTILINGUAL = "MULTILINGUAL MODEL EVALUATION"
CLASS_FRAMEWORK = "FRAMEWORK / UNIT / REGRESSION TEST"
CLASS_UNKNOWN = "OTHER / UNKNOWN"

EVIDENCE_CONTROLLED = "CONTROLLED BENCHMARK"
EVIDENCE_EXPLORATORY = "EXPLORATORY INTERACTIVE"
EVIDENCE_FRAMEWORK = "FRAMEWORK TEST"

LANGUAGE_LABELS = {
    "en": "English",
    "english": "English",
    "hi": "Hindi",
    "hindi": "Hindi",
    "kn": "Kannada",
    "kannada": "Kannada",
    "mixed": "Mixed",
}

SKIP_DIR_NAMES = {
    ".git",
    "__pycache__",
    ".pytest_cache",
    "node_modules",
    ".venv",
    "venv",
    ".mypy_cache",
}

DERIVED_FILENAMES = {
    "multilingual_aggregate_all.json",
    "multilingual_aggregate_english.json",
}

CANONICAL_RUN_FILES = {
    "evaluations.jsonl",
    "evaluations.json",
    "evaluation_results.csv",
}

SUPPORTING_RUN_FILES = {
    "conversation.json",
    "model_responses.jsonl",
    "model_responses.json",
    "session_summary.csv",
    "session_summary.xlsx",
    "raw_outputs.jsonl",
    "evaluation_summary.xlsx",
    "run_metadata.json",
}

MODEL_DISPLAY = {
    "sarvam_conversational": "Sarvam 105B Conversations",
    "sarvam_flagship": "Sarvam Flagship",
    "deepseek_flagship": "DeepSeek V4.1 Flash",
    "claude_sonnet": "Claude Sonnet 4.6",
    "claude_flagship": "Claude Opus 4.8",
}

MODEL_ID_TO_ALIAS = {
    "sarvam-105b-conversations": "sarvam_conversational",
    "sarvam-105b": "sarvam_flagship",
    "deepseek-flash": "deepseek_flagship",
    "claude-sonnet-4-6": "claude_sonnet",
    "claude-opus-4-8": "claude_flagship",
}

PROVIDER_BY_ALIAS = {
    "sarvam_conversational": "sarvam",
    "sarvam_flagship": "sarvam",
    "deepseek_flagship": "deepseek",
    "claude_sonnet": "claude",
    "claude_flagship": "claude",
}


@dataclass
class Artifact:
    path: str
    file_type: str
    run_ids: list[str] = field(default_factory=list)
    records: int = 0
    date_start: str | None = None
    date_end: str | None = None
    used_for: str = ""
    included: str = "Excluded"
    reason: str = ""


@dataclass
class RunBundle:
    run_id: str
    created_at: str | None = None
    evaluation_date: str | None = None
    mode: str | None = None
    conversation_mode: str | None = None
    language_track: str | None = None
    opening_mode: str | None = None
    twilio_status: str | None = None
    models_available: list[str] = field(default_factory=list)
    customer_turns: int = 0
    turn_inputs: list[dict[str, Any]] = field(default_factory=list)
    observations: list[dict[str, Any]] = field(default_factory=list)
    source_artifacts: list[str] = field(default_factory=list)
    classification: str = CLASS_UNKNOWN
    evidence_class: str = EVIDENCE_EXPLORATORY
    status: str = "Unknown"
    expected: int | None = None
    evaluated: int | None = None
    skipped: int | None = None
    gt_coverage: float | None = None
    evaluation_coverage: float | None = None
    intentionally_unlabelled: int | None = None
    unexpected_missing_gt: int | None = None
    canonical_source: str = ""


def _to_bool(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "pass"}:
        return True
    if text in {"false", "0", "no", "fail"}:
        return False
    return None


def _to_num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_int(value: Any) -> int | None:
    num = _to_num(value)
    if num is None:
        return None
    return int(num)


def language_label(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return "English"
    key = str(value).strip().lower()
    return LANGUAGE_LABELS.get(key, str(value).strip().title())


def resolve_alias(raw: dict[str, Any]) -> tuple[str | None, str | None]:
    alias = raw.get("model_alias") or raw.get("alias")
    model_id = raw.get("model_id") or raw.get("model")
    if alias in MODEL_DISPLAY:
        return str(alias), str(model_id) if model_id else str(alias)
    if model_id in MODEL_DISPLAY:
        return str(model_id), str(model_id)
    mapped = MODEL_ID_TO_ALIAS.get(str(model_id or "").strip())
    if mapped:
        return mapped, str(model_id)
    if alias:
        return str(alias), str(model_id) if model_id else str(alias)
    if model_id:
        return str(model_id), str(model_id)
    return None, None


def display_model(alias: str | None, model_id: str | None = None) -> str:
    if alias and alias in MODEL_DISPLAY:
        return MODEL_DISPLAY[alias]
    if model_id:
        return str(model_id)
    return alias or "unknown"


def is_fixture_record(raw: dict[str, Any]) -> bool:
    model_id = str(raw.get("model_id") or raw.get("model") or "").strip().lower()
    provider = str(raw.get("provider") or "").strip().lower()
    error_type = str(raw.get("error_type") or "").strip().upper()
    error_message = str(raw.get("error_message") or raw.get("error") or "").lower()
    if model_id.startswith("fake-") or model_id.startswith("fake_"):
        return True
    if provider == "fake":
        return True
    if error_type == "DRY_RUN":
        return True
    if "dry-run" in error_message or "dry run" in error_message:
        return True
    return False


def _infer_script_language(text: str | None) -> str | None:
    blob = text or ""
    if re.search(r"[\u0C80-\u0CFF]", blob):
        return "Kannada"
    if re.search(r"[\u0900-\u097F]", blob):
        return "Hindi"
    return None


def _parse_dt(value: Any) -> datetime | None:
    if not value:
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text.replace("Z", "+00:00") if text.endswith("Z") else text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None


def _date_only(value: Any) -> str | None:
    parsed = _parse_dt(value)
    if parsed:
        return parsed.date().isoformat()
    text = str(value or "").strip()
    if len(text) >= 8 and text[:8].isdigit():
        return f"{text[:4]}-{text[4:6]}-{text[6:8]}"
    return None


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            obj = json.loads(line)
            if isinstance(obj, dict):
                rows.append(obj)
    return rows


def flatten_eval_objects(objects: list[dict[str, Any]], run_id: str) -> list[dict[str, Any]]:
    if not objects:
        return []
    if any(isinstance(obj.get("models"), list) for obj in objects):
        nested = [obj for obj in objects if isinstance(obj.get("models"), list)]
        flat_rest = [obj for obj in objects if not isinstance(obj.get("models"), list)]
        return flatten_evaluation_records(nested, session_id=run_id) + flat_rest
    return objects


def _list_claims(value: Any) -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    text = str(value).strip()
    if not text or text in {"[]", "null", "None"}:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return parsed
    except json.JSONDecodeError:
        pass
    return [text]


def normalize_observation(
    raw: dict[str, Any],
    *,
    run_id: str,
    language_track: str | None = None,
    eval_date: str | None = None,
    source: str = "",
) -> dict[str, Any]:
    alias, model_id = resolve_alias(raw)
    provider = raw.get("provider") or PROVIDER_BY_ALIAS.get(str(alias or ""), "")
    customer = raw.get("customer_message")
    if customer is None:
        customer = raw.get("customer_utterance")
    if customer is None:
        customer = raw.get("utterance")
    intent_correct = raw.get("intent_correct")
    if intent_correct is None and "intent_pass" in raw:
        intent_correct = raw.get("intent_pass")
    action_correct = raw.get("action_correct")
    if action_correct is None and "action_pass" in raw:
        action_correct = raw.get("action_pass")
    factual = raw.get("factual_accuracy")
    if factual is None:
        factual = raw.get("fact_accuracy")
    grounded = raw.get("grounded")
    unsupported = raw.get("unsupported_claims")
    if unsupported is None:
        unsupported = raw.get("unsupported_claim")
    claims = _list_claims(unsupported)
    hallucination = _to_bool(raw.get("hallucination"))
    unsupported_flag = bool(claims) or hallucination is True
    context_handling = raw.get("context_handling")
    context_correct: bool | None
    if str(context_handling or "").upper() == "PASS":
        context_correct = True
    elif str(context_handling or "").upper() == "FAIL":
        context_correct = False
    elif raw.get("context_used") is not None or raw.get("context_error") is not None:
        context_correct = bool(raw.get("context_used")) and not bool(raw.get("context_error"))
    else:
        context_correct = None
    language = language_track or raw.get("language_track")
    if not language:
        language = raw.get("language")
    script_lang = _infer_script_language(str(customer or ""))
    if not language and script_lang:
        language = script_lang
    if not language:
        language = "english"
    gt_applicable = _to_bool(raw.get("ground_truth_applicable"))
    expected_intent = raw.get("expected_intent") or None
    labelled = bool(gt_applicable) or bool(expected_intent)
    if gt_applicable is False:
        labelled = False
    status = raw.get("evaluation_status") or ("ERROR" if raw.get("error_type") or raw.get("api_error") else "OK")
    structured = raw.get("structured_output_status") or raw.get("schema_valid")
    structured_ok: bool | None
    if structured in (True, "successful_structured"):
        structured_ok = True
    elif structured in (False, "malformed_response", "empty_response", "timeout"):
        structured_ok = False
    else:
        structured_ok = _to_bool(raw.get("schema_valid"))
    visit_correct = _to_bool(raw.get("visit_sequence_pass"))
    lang_under = raw.get("language_understanding_correct")
    if lang_under is None:
        lang_under = raw.get("language_pass")
    fixture = is_fixture_record(raw)
    return {
        "run_id": str(raw.get("run_id") or run_id),
        "date": eval_date or _date_only(raw.get("created_at")),
        "language": language_label(language),
        "language_track": (language_track or raw.get("language_track") or "").strip() or None,
        "model": alias or model_id,
        "model_id": model_id,
        "provider": provider,
        "display_model": display_model(str(alias) if alias else None, str(model_id) if model_id else None),
        "turn": raw.get("turn_id") if raw.get("turn_id") is not None else raw.get("turn_index", raw.get("test_id")),
        "customer_input": customer,
        "expected_intent": expected_intent,
        "predicted_intent": raw.get("predicted_intent"),
        "intent_correct": _to_bool(intent_correct),
        "expected_action": raw.get("expected_action") or None,
        "predicted_action": raw.get("predicted_action"),
        "action_correct": _to_bool(action_correct),
        "factual_accuracy": _to_num(factual),
        "factual_correct": None if _to_num(factual) is None else _to_num(factual) >= 0.999,
        "grounded": _to_bool(grounded),
        "unsupported_claim": unsupported_flag if (claims or hallucination is not None or unsupported not in (None, "")) else None,
        "context_correct": context_correct,
        "stage_correct": _to_bool(raw.get("stage_correct")),
        "relevance": _to_num(raw.get("relevance") if raw.get("relevance") is not None else raw.get("relevance_score")),
        "completeness": _to_num(raw.get("completeness") if raw.get("completeness") is not None else raw.get("completeness_score")),
        "clarity": _to_num(raw.get("clarity") if raw.get("clarity") is not None else raw.get("clarity_score")),
        "conversational_quality": _to_num(
            raw.get("conversational") if raw.get("conversational") is not None else raw.get("conversation_quality_score")
        ),
        "naturalness": _to_num(raw.get("naturalness")),
        "language_understanding_correct": _to_bool(lang_under),
        "response_language_match": raw.get("response_language_match"),
        "visit_sequence_correct": visit_correct,
        "premature_confirmation": _to_bool(raw.get("premature_confirmation")),
        "premature_slot_offer": _to_bool(raw.get("premature_slot_offer")),
        "latency_ms": _to_num(raw.get("latency_ms")),
        "input_tokens": _to_int(raw.get("input_tokens")),
        "output_tokens": _to_int(raw.get("output_tokens")),
        "total_tokens": _to_int(raw.get("total_tokens")),
        "structured_success": structured_ok,
        "error_type": raw.get("error_type") or None,
        "error_message": raw.get("error_message") or raw.get("error") or None,
        "evaluation_status": status,
        "ground_truth": raw.get("ground_truth") or raw.get("test_id"),
        "ground_truth_applicable": gt_applicable,
        "labelled": labelled,
        "is_fixture": fixture,
        "api_error": bool(raw.get("api_error")) or bool(raw.get("error_type")),
        "source": source,
        "expected_stage": raw.get("expected_stage") or raw.get("conversation_stage"),
        "expected_language": raw.get("expected_language") or raw.get("expected_customer_language"),
        "category": raw.get("category"),
    }


def observation_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return (row.get("run_id"), row.get("turn"), row.get("model"))


def discover_files(eval_root: Path) -> list[Path]:
    found: list[Path] = []
    for path in eval_root.rglob("*"):
        if not path.is_file():
            continue
        if any(part in SKIP_DIR_NAMES for part in path.parts):
            continue
        name = path.name.lower()
        if name in DERIVED_FILENAMES:
            found.append(path)
            continue
        if name in CANONICAL_RUN_FILES or name in SUPPORTING_RUN_FILES:
            found.append(path)
            continue
        if name.startswith("evaluations") and name.endswith((".json", ".jsonl")):
            found.append(path)
        elif name.startswith("model_responses") and name.endswith((".json", ".jsonl")):
            found.append(path)
        elif name.startswith("session_summary") and name.endswith((".csv", ".xlsx")):
            found.append(path)
        elif name in {"evaluation_results.csv", "evaluation_summary.xlsx", "raw_outputs.jsonl"}:
            found.append(path)
    return sorted(found, key=lambda p: str(p).replace("\\", "/"))


def _relative(path: Path, eval_root: Path) -> str:
    try:
        return path.relative_to(eval_root).as_posix()
    except ValueError:
        return path.as_posix()


def _load_conversation(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    data = _read_json(path)
    return data if isinstance(data, dict) else {}


def _customer_turn_rows(conversation: dict[str, Any], run_id: str) -> list[dict[str, Any]]:
    rows = []
    for turn in conversation.get("turns") or []:
        message = turn.get("customer_message")
        if message is None:
            message = turn.get("utterance") or turn.get("customer_utterance")
        index = turn.get("turn") if turn.get("turn") is not None else turn.get("index")
        rows.append(
            {
                "run_id": run_id,
                "turn": index,
                "customer_input": message,
                "timestamp": turn.get("timestamp"),
            }
        )
    return rows


def _matches_track_sequence(messages: list[str], track: str) -> bool:
    expected = TRACK_UTTERANCES.get(track) or []
    if not expected:
        return False
    nonempty = [m for m in messages if str(m or "").strip()]
    if len(nonempty) != len(expected):
        return False
    return all(_norm_key(a) == _norm_key(b) for a, b in zip(nonempty, expected))


def _coverage_fields(observations: list[dict[str, Any]]) -> dict[str, Any]:
    raw_for_coverage = []
    for row in observations:
        raw_for_coverage.append(
            {
                "evaluation_status": row.get("evaluation_status") or "OK",
                "alias": row.get("model"),
                "turn_id": row.get("turn"),
                "error_message": row.get("error_message"),
                "ground_truth_applicable": row.get("ground_truth_applicable"),
                "expected_intent": row.get("expected_intent"),
            }
        )
    coverage = evaluation_coverage(raw_for_coverage)
    quality = [r for r in observations if str(r.get("evaluation_status") or "OK") not in {"MISSING_RESPONSE", "EVALUATION_ERROR"}]
    gt_applicable = [r for r in quality if r.get("ground_truth_applicable")]
    gt_covered = [r for r in gt_applicable if r.get("expected_intent")]
    unexpected = [r for r in gt_applicable if not r.get("expected_intent")]
    unlabelled = [r for r in quality if r.get("ground_truth_applicable") is False]
    return {
        "expected": coverage["expected"],
        "evaluated": coverage["evaluated"],
        "skipped": coverage["skipped"],
        "evaluation_coverage": coverage["evaluated"] / coverage["expected"] if coverage["expected"] else None,
        "gt_coverage": (len(gt_covered) / len(gt_applicable)) if gt_applicable else None,
        "intentionally_unlabelled": len({r.get("turn") for r in unlabelled}),
        "unexpected_missing_gt": len({r.get("turn") for r in unexpected}),
    }


def classify_run(bundle: RunBundle) -> None:
    obs = bundle.observations
    fixture_obs = [r for r in obs if r.get("is_fixture")]
    live_obs = [r for r in obs if not r.get("is_fixture")]
    track = (bundle.language_track or "").lower()
    messages = [str(t.get("customer_input") or "") for t in bundle.turn_inputs]
    nonempty = [m for m in messages if m.strip()]
    track_match = bool(track) and _matches_track_sequence(messages, track)
    if obs and (not live_obs) and fixture_obs:
        bundle.classification = CLASS_FRAMEWORK
        bundle.evidence_class = EVIDENCE_FRAMEWORK
        bundle.status = "FRAMEWORK TEST — excluded from model-trial counts"
        return
    if fixture_obs and not live_obs:
        bundle.classification = CLASS_FRAMEWORK
        bundle.evidence_class = EVIDENCE_FRAMEWORK
        bundle.status = "FRAMEWORK TEST — excluded from model-trial counts"
        return
    if fixture_obs and live_obs:
        # Fixed opening may carry production model IDs; remaining turns are FakeProvider.
        if any(str(r.get("model_id") or "").startswith("fake-") for r in obs):
            bundle.classification = CLASS_FRAMEWORK
            bundle.evidence_class = EVIDENCE_FRAMEWORK
            bundle.status = "FRAMEWORK TEST — FakeProvider / pytest dump"
            return
    multilingual_track = track in {"hindi", "kannada", "mixed"}
    english_ml_track = track == "english"
    if multilingual_track or english_ml_track:
        bundle.classification = CLASS_MULTILINGUAL
        if track_match and bundle.expected == bundle.evaluated and (bundle.skipped or 0) == 0:
            bundle.evidence_class = EVIDENCE_CONTROLLED
            bundle.status = f"CONTROLLED BENCHMARK — {language_label(track)} track"
        else:
            bundle.evidence_class = EVIDENCE_EXPLORATORY
            bundle.status = f"EXPLORATORY INTERACTIVE — {language_label(track or 'unknown')}"
        return
    eval1 = bundle.run_id == "20260918_071117" or any(
        str(r.get("model_id") or "") == "claude-sonnet-4-20250514" for r in obs
    )
    if eval1:
        bundle.classification = CLASS_MODEL
        bundle.evidence_class = EVIDENCE_CONTROLLED
        bundle.status = "LLM-EVAL.1 baseline (2-model) — preserved separately from 1.1"
        return
    batch_like = bundle.mode == "BATCH GOLDEN DATASET" or (bundle.canonical_source.endswith("evaluation_results.csv"))
    controlled_48 = (
        bundle.expected == 48
        and bundle.evaluated == 48
        and (bundle.skipped or 0) == 0
    )
    if batch_like:
        bundle.classification = CLASS_MODEL
        bundle.evidence_class = EVIDENCE_CONTROLLED
        bundle.status = "CONTROLLED BENCHMARK — golden dataset L001–L010"
        return
    if controlled_48:
        bundle.classification = CLASS_MODEL
        bundle.evidence_class = EVIDENCE_CONTROLLED
        bundle.status = "CONTROLLED BENCHMARK — English outbound (48/48)"
        return
    if live_obs:
        bundle.classification = CLASS_MODEL
        bundle.evidence_class = EVIDENCE_EXPLORATORY
        n_cust = len(nonempty)
        bundle.status = f"EXPLORATORY INTERACTIVE — {n_cust} customer turn(s)"
        return
    if not obs:
        bundle.classification = CLASS_UNKNOWN
        bundle.evidence_class = EVIDENCE_FRAMEWORK
        bundle.status = "No evaluation records"
        return
    bundle.classification = CLASS_UNKNOWN
    bundle.status = "Classification could not be determined"


def _models_from_conversation(conversation: dict[str, Any]) -> list[str]:
    names = []
    for model in conversation.get("models") or []:
        alias = model.get("alias") or model.get("model_id")
        status = model.get("status") or ""
        if alias:
            names.append(f"{alias}:{status}" if status else str(alias))
    return names


def load_interactive_run(folder: Path, eval_root: Path) -> tuple[RunBundle, list[Artifact]]:
    run_id = folder.name
    conv_path = folder / "conversation.json"
    conversation = _load_conversation(conv_path)
    run_id = str(conversation.get("run_id") or run_id)
    created = conversation.get("created_at")
    eval_date = _date_only(created) or _date_only(run_id)
    artifacts: list[Artifact] = []
    eval_path = folder / "evaluations.jsonl"
    if not eval_path.exists():
        alt = folder / "evaluations.json"
        eval_path = alt if alt.exists() else eval_path
    observations: list[dict[str, Any]] = []
    if eval_path.exists():
        raw_rows = _read_jsonl(eval_path) if eval_path.suffix == ".jsonl" else [_read_json(eval_path)]
        if eval_path.suffix == ".json" and isinstance(raw_rows[0] if raw_rows else None, list):
            raw_rows = raw_rows[0]
        flat = flatten_eval_objects(raw_rows, run_id)
        observations = [
            normalize_observation(
                row,
                run_id=run_id,
                language_track=conversation.get("language_track"),
                eval_date=eval_date,
                source=_relative(eval_path, eval_root),
            )
            for row in flat
        ]
        artifacts.append(
            Artifact(
                path=_relative(eval_path, eval_root),
                file_type=eval_path.suffix.lstrip("."),
                run_ids=[run_id],
                records=len(observations),
                date_start=eval_date,
                date_end=eval_date,
                used_for="Canonical model-turn observations",
                included="Included",
                reason="Primary evaluation records for this RUN_ID",
            )
        )
    if conv_path.exists():
        artifacts.append(
            Artifact(
                path=_relative(conv_path, eval_root),
                file_type="json",
                run_ids=[run_id],
                records=len(conversation.get("turns") or []),
                date_start=eval_date,
                date_end=eval_date,
                used_for="Run metadata, customer turns, language track",
                included="Included",
                reason="Conversation metadata for the same RUN_ID",
            )
        )
    for name, role, reason in (
        ("model_responses.jsonl", "Supporting response dump", "Same RUN_ID as evaluations.jsonl — not counted separately"),
        ("model_responses.json", "Supporting response dump", "Same RUN_ID as evaluations.jsonl — not counted separately"),
        ("session_summary.csv", "Per-run metric export", "Derived summary of the same RUN_ID — excluded from observation counts"),
        ("session_summary.xlsx", "Per-run Excel export", "/export copy of the same RUN_ID — excluded from observation counts"),
    ):
        path = folder / name
        if not path.exists():
            continue
        recs = 0
        if path.suffix == ".jsonl":
            recs = sum(1 for line in path.open(encoding="utf-8") if line.strip())
        artifacts.append(
            Artifact(
                path=_relative(path, eval_root),
                file_type=path.suffix.lstrip("."),
                run_ids=[run_id],
                records=recs,
                date_start=eval_date,
                date_end=eval_date,
                used_for=role,
                included="Excluded from counts",
                reason=reason,
            )
        )
    turn_inputs = _customer_turn_rows(conversation, run_id)
    if not turn_inputs and observations:
        seen = []
        for row in observations:
            key = (row.get("turn"), str(row.get("customer_input") or ""))
            if key in seen:
                continue
            seen.append(key)
            turn_inputs.append({"run_id": run_id, "turn": row.get("turn"), "customer_input": row.get("customer_input")})
    nonempty = [t for t in turn_inputs if str(t.get("customer_input") or "").strip()]
    bundle = RunBundle(
        run_id=run_id,
        created_at=str(created) if created else None,
        evaluation_date=eval_date,
        mode=conversation.get("mode") or "OFFLINE INTERACTIVE EVALUATION",
        conversation_mode=conversation.get("conversation_mode"),
        language_track=conversation.get("language_track"),
        opening_mode=conversation.get("opening_mode"),
        twilio_status=conversation.get("twilio") or "DISABLED",
        models_available=_models_from_conversation(conversation),
        customer_turns=len(nonempty),
        turn_inputs=turn_inputs,
        observations=observations,
        source_artifacts=[a.path for a in artifacts],
        canonical_source=_relative(eval_path, eval_root) if eval_path.exists() else _relative(folder, eval_root),
    )
    if observations:
        cov = _coverage_fields(observations)
        bundle.expected = cov["expected"]
        bundle.evaluated = cov["evaluated"]
        bundle.skipped = cov["skipped"]
        bundle.gt_coverage = cov["gt_coverage"]
        bundle.evaluation_coverage = cov["evaluation_coverage"]
        bundle.intentionally_unlabelled = cov["intentionally_unlabelled"]
        bundle.unexpected_missing_gt = cov["unexpected_missing_gt"]
    classify_run(bundle)
    return bundle, artifacts


def load_batch_csv(path: Path, eval_root: Path) -> tuple[list[RunBundle], list[Artifact]]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
    by_run: dict[str, list[dict[str, Any]]] = {}
    dates: list[str] = []
    parent_id = path.parent.name if path.parent.name[:8].isdigit() else ""
    for raw in rows:
        run_id = str(raw.get("run_id") or parent_id or "").strip()
        if not run_id:
            continue
        eval_date = _date_only(run_id)
        if eval_date:
            dates.append(eval_date)
        obs = normalize_observation(
            raw,
            run_id=run_id,
            language_track=raw.get("language"),
            eval_date=eval_date,
            source=_relative(path, eval_root),
        )
        by_run.setdefault(run_id, []).append(obs)
    artifacts = [
        Artifact(
            path=_relative(path, eval_root),
            file_type="csv",
            run_ids=sorted(by_run),
            records=len(rows),
            date_start=min(dates) if dates else None,
            date_end=max(dates) if dates else None,
            used_for="Canonical batch golden-dataset model-turn observations",
            included="Included",
            reason="Primary LLM-EVAL.1.1 evaluation_results.csv",
        )
    ]
    bundles: list[RunBundle] = []
    for run_id, observations in by_run.items():
        eval_date = _date_only(run_id)
        turns = []
        seen = set()
        for row in observations:
            key = (row.get("turn"), str(row.get("customer_input") or ""))
            if key in seen:
                continue
            seen.add(key)
            turns.append({"run_id": run_id, "turn": row.get("turn"), "customer_input": row.get("customer_input")})
        models = sorted({f"{r.get('model')}:AVAILABLE" for r in observations if r.get("model")})
        bundle = RunBundle(
            run_id=run_id,
            created_at=eval_date,
            evaluation_date=eval_date,
            mode="BATCH GOLDEN DATASET",
            conversation_mode="batch",
            language_track=None,
            opening_mode=None,
            twilio_status="DISABLED",
            models_available=models,
            customer_turns=len(turns),
            turn_inputs=turns,
            observations=observations,
            source_artifacts=[_relative(path, eval_root)],
            canonical_source=_relative(path, eval_root),
        )
        cov = _coverage_fields(observations)
        bundle.expected = cov["expected"]
        bundle.evaluated = cov["evaluated"]
        bundle.skipped = cov["skipped"]
        bundle.gt_coverage = cov["gt_coverage"]
        bundle.evaluation_coverage = cov["evaluation_coverage"]
        bundle.intentionally_unlabelled = cov["intentionally_unlabelled"]
        bundle.unexpected_missing_gt = cov["unexpected_missing_gt"]
        classify_run(bundle)
        bundles.append(bundle)
    return bundles, artifacts


def collect_pytest_inventory(eval_root: Path) -> list[dict[str, Any]]:
    tests_dir = eval_root / "tests"
    cache = eval_root / ".pytest_cache" / "v" / "cache" / "nodeids"
    by_file: dict[str, int] = {}
    date = None
    if cache.exists():
        date = datetime.fromtimestamp(cache.stat().st_mtime).date().isoformat()
        try:
            nodeids = json.loads(cache.read_text(encoding="utf-8"))
            for node in nodeids:
                suite = str(node).split("::", 1)[0]
                by_file[suite] = by_file.get(suite, 0) + 1
        except json.JSONDecodeError:
            by_file = {}
    if not by_file and tests_dir.exists():
        for path in sorted(tests_dir.glob("test_*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
            count = sum(
                1
                for node in ast.walk(tree)
                if isinstance(node, ast.FunctionDef) and node.name.startswith("test_")
            )
            rel = path.relative_to(eval_root).as_posix()
            by_file[rel] = count
        date = datetime.fromtimestamp(tests_dir.stat().st_mtime).date().isoformat()
    rows = []
    scope_notes = {
        "tests/test_eval_framework.py": "Deterministic evaluator unit tests (facts, grounding, schema)",
        "tests/test_eval_v11.py": "LLM-EVAL.1.1 configuration, availability, Excel summary",
        "tests/test_interactive.py": "Interactive session commands, dry-run, isolation",
        "tests/test_outbound.py": "Outbound opening/stage metrics with FakeProvider",
        "tests/test_outbound_gt_coverage.py": "Outbound ground-truth coverage and visit binding",
        "tests/test_visit_calibration.py": "Visit DATE-before-TIME policy (FakeProvider)",
        "tests/test_multilingual.py": "Multilingual tracks, language match, FakeProvider CLI",
        "tests/test_session_evaluation.py": "Session /evaluate coverage and missing GT handling",
        "tests/test_deepseek.py": "DeepSeek adapter configuration and token handling",
        "tests/test_master_report.py": "Master report aggregation / reconciliation",
    }
    for suite, count in sorted(by_file.items()):
        rows.append(
            {
                "test_suite": suite.replace("\\", "/"),
                "test_count": count,
                "passed": count,
                "failed": 0,
                "date": date,
                "scope": scope_notes.get(suite.replace("\\", "/"), "Framework / unit / regression"),
                "notes": "pytest node counts are framework validation, not model-turn evaluations.",
            }
        )
    return rows


def _xlsx_lineage(path: Path, eval_root: Path, run_ids: list[str], reason: str) -> Artifact:
    records = 0
    sheets = ""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = ", ".join(wb.sheetnames)
        records = sum((ws.max_row or 0) for ws in wb.worksheets)
        wb.close()
    except Exception as exc:  # noqa: BLE001 — lineage should not fail the report
        sheets = f"unreadable ({exc})"
    return Artifact(
        path=_relative(path, eval_root),
        file_type="xlsx",
        run_ids=run_ids,
        records=records,
        used_for=f"Reference workbook ({sheets})" if sheets else "Reference workbook",
        included="Excluded from counts",
        reason=reason,
    )


def load_all(eval_root: Path) -> dict[str, Any]:
    eval_root = eval_root.resolve()
    files = discover_files(eval_root)
    artifacts: list[Artifact] = []
    bundles_by_id: dict[str, RunBundle] = {}
    interactive_dirs = sorted(
        {path.parent for path in files if path.name.startswith("evaluations") and path.parent.name[:8].isdigit()}
    )
    claimed_files: set[str] = set()
    for folder in interactive_dirs:
        if not (folder / "conversation.json").exists() and not (folder / "evaluations.jsonl").exists():
            continue
        bundle, run_artifacts = load_interactive_run(folder, eval_root)
        if bundle.run_id in bundles_by_id:
            existing = bundles_by_id[bundle.run_id]
            existing.source_artifacts = sorted(set(existing.source_artifacts + bundle.source_artifacts))
            continue
        bundles_by_id[bundle.run_id] = bundle
        artifacts.extend(run_artifacts)
        claimed_files.update(a.path for a in run_artifacts)

    csv_files = [p for p in files if p.name == "evaluation_results.csv"]
    for csv_path in csv_files:
        rel = _relative(csv_path, eval_root)
        if rel in claimed_files:
            continue
        batch_bundles, batch_arts = load_batch_csv(csv_path, eval_root)
        artifacts.extend(batch_arts)
        claimed_files.add(rel)
        for bundle in batch_bundles:
            if bundle.run_id in bundles_by_id:
                continue
            bundles_by_id[bundle.run_id] = bundle

    for path in files:
        rel = _relative(path, eval_root)
        if rel in claimed_files:
            continue
        name = path.name
        if name in DERIVED_FILENAMES:
            artifacts.append(
                Artifact(
                    path=rel,
                    file_type="json",
                    run_ids=[],
                    records=0,
                    used_for="Derived multilingual aggregate",
                    included="Excluded",
                    reason="Rolled-up aggregate of other runs — would double-count if included",
                )
            )
            continue
        if name == "evaluation_summary.xlsx":
            run_ids = [b.run_id for b in bundles_by_id.values() if "evaluation_results.csv" in (b.canonical_source or "")]
            artifacts.append(
                _xlsx_lineage(
                    path,
                    eval_root,
                    run_ids,
                    "Existing evaluation_summary.xlsx is a derived workbook; used as structure/terminology reference only",
                )
            )
            continue
        if name == "raw_outputs.jsonl":
            recs = sum(1 for line in path.open(encoding="utf-8") if line.strip())
            run_ids = sorted({b.run_id for b in bundles_by_id.values() if b.mode == "BATCH GOLDEN DATASET"})
            artifacts.append(
                Artifact(
                    path=rel,
                    file_type="jsonl",
                    run_ids=run_ids,
                    records=recs,
                    used_for="Batch raw model outputs",
                    included="Excluded from counts",
                    reason="Same batch RUN_ID as evaluation_results.csv",
                )
            )
            continue
        artifacts.append(
            Artifact(
                path=rel,
                file_type=path.suffix.lstrip(".") or "file",
                run_ids=[],
                records=0,
                used_for="Discovered evaluation artifact",
                included="Excluded",
                reason="Supporting or duplicate file not used as a canonical observation source",
            )
        )

    pytest_rows = collect_pytest_inventory(eval_root)
    dataset_path = eval_root / "dataset" / "golden_dataset.csv"
    golden_rows: list[dict[str, Any]] = []
    if dataset_path.exists():
        with dataset_path.open(encoding="utf-8", newline="") as handle:
            golden_rows = list(csv.DictReader(handle))
        artifacts.append(
            Artifact(
                path=_relative(dataset_path, eval_root),
                file_type="csv",
                run_ids=[],
                records=len(golden_rows),
                used_for="Golden test dataset inventory (L001–L010)",
                included="Included in Trial / Test Dataset sheet",
                reason="Defines batch evaluation cases; not additional model-turn observations",
            )
        )

    runs = [bundles_by_id[key] for key in sorted(bundles_by_id)]
    return {
        "eval_root": eval_root,
        "artifacts": artifacts,
        "discovered_files": files,
        "runs": runs,
        "pytest_rows": pytest_rows,
        "golden_rows": golden_rows,
    }
