from __future__ import annotations

import csv
import json
import logging
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from evaluator.dataset import load_dataset
from evaluator.runner import (
    load_env,
    load_text,
    load_yaml_config,
    new_run_id,
    root_dir,
    select_models,
)
from models import build_provider
from models.availability import check_claude_model, check_deepseek_model, check_sarvam_model
from models.prompting import build_user_payload
from openpyxl import Workbook
from openpyxl.styles import Font

from .conversation import SHARED_PROJECT_CONTEXT
from .evaluator import (
    EVAL_STATUS_EVALUATION_ERROR,
    evaluation_coverage,
    evaluate_turn,
    flatten_evaluation_records,
    session_summary,
)
from .outbound import (
    FIXED_OPENING,
    OPENING_INSTRUCTION,
    evaluate_opening,
    fixed_opening_responses,
    shared_stage,
)
from .visit_policy import SITE_VISIT_SCHEDULING_RULE
from .renderer import redact
from .schemas import DISPLAY_ORDER, DISPLAY_NAMES, ModelSlot, TurnResponse

LOGGER = logging.getLogger("llm_eval.interactive")


def bootstrap_path() -> Path:
    eval_root = Path(__file__).resolve().parents[1]
    if str(eval_root) not in sys.path:
        sys.path.insert(0, str(eval_root))
    return eval_root


def _opening_metric_rows(evaluations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    opening = next((item for item in evaluations if item.get("kind") == "opening"), None)
    if not opening:
        return []
    fields = (
        "opening_element_company",
        "opening_element_project",
        "opening_element_description",
        "opening_element_location",
        "opening_element_permission",
        "opening_conciseness",
        "outbound_appropriateness",
    )
    by_alias = {model.get("alias"): model for model in opening.get("models") or []}
    rows = []
    for field in fields:
        rows.append({"metric": field, **{alias: (by_alias.get(alias) or {}).get(field) for alias in DISPLAY_ORDER}})
    return rows


def _safe_error(message: str | None) -> str | None:
    if not message:
        return None
    return redact(message)


class InteractiveSession:
    def __init__(
        self,
        dry_run: bool = False,
        max_display_chars: int = 0,
        config_path: Path | None = None,
        providers: dict[str, Any] | None = None,
        conversation_mode: str = "customer",
        opening_mode: str = "generated",
    ):
        bootstrap_path()
        load_env()
        self.dry_run = dry_run
        self.max_display_chars = max_display_chars
        self.conversation_mode = "outbound" if conversation_mode == "outbound" else "customer"
        self.opening_mode = "fixed" if opening_mode == "fixed" else "generated"
        self.eval_root = root_dir()
        self.config = load_yaml_config(config_path)
        self.system_prompt = load_text(self.eval_root / "prompts" / "calling_agent_system_prompt.txt")
        if self.conversation_mode == "outbound":
            self.system_prompt = self.system_prompt.rstrip() + "\n\n" + SITE_VISIT_SCHEDULING_RULE
        self.retrieved_context = SHARED_PROJECT_CONTEXT
        dataset_path = self.eval_root / "dataset" / "golden_dataset.csv"
        self.dataset = load_dataset(str(dataset_path))
        selected, skipped = select_models(self.config, list(DISPLAY_ORDER))
        self.skip_notes = skipped
        by_name = {item["name"]: item for item in selected}
        self.slots: list[ModelSlot] = []
        self._runtimes: dict[str, dict[str, Any]] = {}
        self._providers: dict[str, Any] = providers or {}
        for alias in DISPLAY_ORDER:
            runtime = by_name.get(alias)
            if runtime is None:
                self.slots.append(
                    ModelSlot(
                        alias=alias,
                        provider=alias.split("_")[0],
                        role="",
                        model_id="",
                        available=False,
                        status="UNAVAILABLE",
                        error=f"not in active configuration ({'; '.join(skipped) or 'unknown'})",
                    )
                )
                continue
            available = bool(runtime.get("available"))
            error = None
            status = "AVAILABLE" if available else "UNAVAILABLE"
            if not runtime.get("api_key"):
                available = False
                error = f"{runtime.get('api_key_env')}: NOT CONFIGURED"
                status = "UNAVAILABLE — missing API key"
            elif not runtime.get("model"):
                available = False
                error = f"missing model id ({runtime.get('model_env')})"
                status = "UNAVAILABLE"
            slot = ModelSlot(
                alias=alias,
                provider=runtime["provider"],
                role=runtime.get("role") or "",
                model_id=runtime.get("model") or "",
                available=available,
                status=status,
                error=error,
                base_url=runtime.get("base_url") or "",
                runtime=runtime,
            )
            self.slots.append(slot)
            self._runtimes[alias] = runtime
            if providers is None and available and not dry_run:
                provider_kwargs = {
                    "api_key": runtime["api_key"],
                    "model": runtime["model"],
                    "base_url": runtime["base_url"],
                    "timeout": runtime["timeout"],
                    "temperature": runtime["temperature"],
                    "max_tokens": runtime["max_tokens"],
                    "max_retries": runtime["max_retries"],
                    "model_alias": alias,
                    "model_role": runtime.get("role"),
                }
                if runtime.get("thinking_mode"):
                    provider_kwargs["thinking_mode"] = runtime["thinking_mode"]
                self._providers[alias] = build_provider(runtime["provider"], **provider_kwargs)
        self.run_id = new_run_id()
        self.created_at = datetime.now(timezone.utc).isoformat()
        self.histories: dict[str, list[dict[str, str]]] = {alias: [] for alias in DISPLAY_ORDER}
        self.turns: list[dict[str, Any]] = []
        self.evaluations: list[dict[str, Any]] = []
        self.last_debug: dict[str, Any] = {}
        self.exported_dir: Path | None = None

    def reset(self) -> None:
        self.histories = {alias: [] for alias in DISPLAY_ORDER}
        self.turns = []
        self.evaluations = []
        self.last_debug = {}

    def next_turn_index(self) -> int:
        return len(self.turns) + 1

    def generate_turn(self, utterance: str) -> list[TurnResponse]:
        turn_index = self.next_turn_index()
        customer_history = [t["utterance"] for t in self.turns if t.get("kind") != "opening"]
        stage = shared_stage(utterance, turn_index, self.conversation_mode, customer_history)
        extra = {"conversation_stage": stage, "stage": stage}
        histories_snapshot = {alias: list(items) for alias, items in self.histories.items()}
        payloads = {}
        for alias in DISPLAY_ORDER:
            payloads[alias] = {
                "system_prompt": self.system_prompt,
                "conversation_history": histories_snapshot[alias],
                "customer_utterance": utterance,
                "retrieved_context": self.retrieved_context,
                "extra": extra,
            }
        self.last_debug = {
            "user_message": utterance,
            "retrieved_knowledge": self.retrieved_context,
            "system_instructions_summary": self.system_prompt.split("\n", 1)[0][:240],
            "conversation_mode": self.conversation_mode,
            "opening_mode": self.opening_mode if self.conversation_mode == "outbound" else None,
            "conversation_stage": extra["conversation_stage"],
            "same_user_message": True,
            "same_retrieved_knowledge": True,
            "same_system_prompt": True,
            "per_model_history": {alias: list(histories_snapshot[alias]) for alias in DISPLAY_ORDER},
            "model_request_metadata": {
                alias: {
                    "provider": slot.provider,
                    "model_id": slot.model_id,
                    "status": slot.status,
                    "max_tokens": (slot.runtime or {}).get("max_tokens"),
                    "timeout": (slot.runtime or {}).get("timeout"),
                    "thinking_mode": (slot.runtime or {}).get("thinking_mode"),
                }
                for slot in self.slots
                for alias in [slot.alias]
            },
        }
        ordered = self._collect(payloads)
        for resp in ordered:
            self.histories[resp.alias].append({"role": "user", "content": utterance})
            if not resp.error_type:
                self.histories[resp.alias].append(
                    {"role": "assistant", "content": resp.answer or resp.raw_response or ""}
                )
        self.turns.append(
            {
                "index": turn_index,
                "timestamp": datetime.now(timezone.utc).isoformat(),
                "utterance": utterance,
                "responses": [asdict(item) for item in ordered],
                "histories_before": histories_snapshot,
                "kind": "customer",
            }
        )
        return ordered

    def begin_outbound(self) -> list[TurnResponse]:
        """AI initiates. Generated openings are independent. Fixed openings are identical and offline."""
        if self.conversation_mode != "outbound":
            raise ValueError("begin_outbound requires conversation_mode='outbound'")
        histories_snapshot = {alias: list(items) for alias, items in self.histories.items()}
        if self.opening_mode == "fixed":
            ordered = fixed_opening_responses(self.slots)
            instruction = None
        else:
            instruction = OPENING_INSTRUCTION
            extra = {"conversation_stage": "greeting", "stage": "greeting"}
            payloads = {
                alias: {
                    "system_prompt": self.system_prompt,
                    "conversation_history": histories_snapshot[alias],
                    "customer_utterance": instruction,
                    "retrieved_context": self.retrieved_context,
                    "extra": extra,
                }
                for alias in DISPLAY_ORDER
            }
            ordered = self._collect(payloads)
        for resp in ordered:
            if not resp.error_type:
                self.histories[resp.alias].append(
                    {"role": "assistant", "content": resp.answer or resp.raw_response or ""}
                )
        self.turns.append(
            {
                "index": len(self.turns) + 1,
                "timestamp": datetime.now(timezone.utc).isoformat(),
                "utterance": "",
                "kind": "opening",
                "opening_mode": self.opening_mode,
                "instruction": instruction,
                "responses": [asdict(item) for item in ordered],
                "histories_before": histories_snapshot,
            }
        )
        self.last_debug = {
            "user_message": instruction or FIXED_OPENING,
            "opening_mode": self.opening_mode,
            "conversation_mode": self.conversation_mode,
            "retrieved_knowledge": self.retrieved_context,
            "system_instructions_summary": self.system_prompt.split("\n", 1)[0][:240],
            "conversation_stage": "greeting",
            "same_user_message": True,
            "same_retrieved_knowledge": True,
            "same_system_prompt": True,
            "per_model_history": {alias: list(histories_snapshot[alias]) for alias in DISPLAY_ORDER},
            "model_request_metadata": {
                slot.alias: {
                    "provider": slot.provider,
                    "model_id": slot.model_id,
                    "status": slot.status,
                    "thinking_mode": (slot.runtime or {}).get("thinking_mode"),
                }
                for slot in self.slots
            },
        }
        return ordered

    def _collect(self, payloads: dict[str, dict[str, Any]]) -> list[TurnResponse]:
        responses: dict[str, TurnResponse] = {}
        if self.dry_run:
            for slot in self.slots:
                responses[slot.alias] = TurnResponse(
                    alias=slot.alias,
                    provider=slot.provider,
                    model_id=slot.model_id,
                    answer=None,
                    intent=None,
                    action=None,
                    language=None,
                    stage=None,
                    raw_response="",
                    latency_ms=None,
                    input_tokens=None,
                    output_tokens=None,
                    total_tokens=None,
                    error_type="DRY_RUN",
                    error_message="Dry-run: no API call made",
                    schema_valid=None,
                )
        else:
            with ThreadPoolExecutor(max_workers=4) as pool:
                futures = {
                    pool.submit(self._call_model, slot, payloads[slot.alias]): slot.alias for slot in self.slots
                }
                for future in as_completed(futures):
                    alias = futures[future]
                    try:
                        responses[alias] = future.result()
                    except Exception as exc:
                        slot = next(s for s in self.slots if s.alias == alias)
                        responses[alias] = TurnResponse(
                            alias=alias,
                            provider=slot.provider,
                            model_id=slot.model_id,
                            answer=None,
                            intent=None,
                            action=None,
                            language=None,
                            stage=None,
                            raw_response="",
                            latency_ms=None,
                            input_tokens=None,
                            output_tokens=None,
                            total_tokens=None,
                            error_type="API_ERROR",
                            error_message=_safe_error(str(exc)),
                            schema_valid=None,
                        )
        return [responses[alias] for alias in DISPLAY_ORDER]

    def _call_model(self, slot: ModelSlot, payload: dict[str, Any]) -> TurnResponse:
        if slot.alias not in self._providers:
            return TurnResponse(
                alias=slot.alias,
                provider=slot.provider,
                model_id=slot.model_id,
                answer=None,
                intent=None,
                action=None,
                language=None,
                stage=None,
                raw_response="",
                latency_ms=None,
                input_tokens=None,
                output_tokens=None,
                total_tokens=None,
                error_type="UNAVAILABLE",
                error_message=_safe_error(slot.error or "model unavailable"),
                schema_valid=None,
            )
        result = self._providers[slot.alias].generate(**payload)
        usage = result.get("usage") or {}
        error_type = result.get("error_type")
        return TurnResponse(
            alias=slot.alias,
            provider=slot.provider,
            model_id=result.get("model_id") or slot.model_id,
            answer=result.get("answer"),
            intent=result.get("intent"),
            action=result.get("action"),
            language=result.get("language"),
            stage=None,
            raw_response=result.get("raw_response") or "",
            latency_ms=result.get("latency_ms"),
            input_tokens=usage.get("input_tokens"),
            output_tokens=usage.get("output_tokens"),
            total_tokens=usage.get("total_tokens"),
            error_type=error_type,
            error_message=_safe_error(result.get("error_message") or result.get("error")),
            schema_valid=result.get("schema_valid"),
            confidence=result.get("confidence"),
        )

    def evaluate_latest(self) -> dict[str, Any] | None:
        if not self.turns:
            return None
        ev = self._evaluate_stored_turn(self.turns[-1])
        self.evaluations = [item for item in self.evaluations if item.get("turn_index") != ev.get("turn_index")]
        self.evaluations.append(ev)
        return ev

    def evaluate_session(self) -> dict[str, Any]:
        """Evaluate every stored turn × every configured model. Does not call live models."""
        n_models = len(DISPLAY_ORDER)
        expected = len(self.turns) * n_models
        evaluations: list[dict[str, Any]] = []
        for turn in self.turns:
            try:
                ev = self._evaluate_stored_turn(turn)
            except Exception as exc:
                ev = self._error_turn_evaluation(turn, exc)
            evaluations.append(ev)
        self.evaluations = evaluations
        records = flatten_evaluation_records(evaluations, session_id=self.run_id)
        coverage = evaluation_coverage(records, expected=expected)
        coverage["n_turns"] = len(self.turns)
        coverage["n_models"] = n_models
        return coverage

    def _evaluate_stored_turn(self, turn: dict[str, Any]) -> dict[str, Any]:
        responses = [TurnResponse(**item) for item in turn.get("responses") or []]
        prior = turn.get("histories_before") or {alias: [] for alias in DISPLAY_ORDER}
        if turn.get("kind") == "opening":
            ev = evaluate_opening(
                responses,
                self.retrieved_context,
                turn.get("opening_mode") or self.opening_mode,
                turn["index"],
                required_aliases=list(DISPLAY_ORDER),
            )
        else:
            ev = evaluate_turn(
                turn["utterance"],
                responses,
                self.dataset,
                prior,
                turn["index"],
                self.retrieved_context,
                outbound=self.conversation_mode == "outbound",
                required_aliases=list(DISPLAY_ORDER),
            )
        return self._attach_record_identity(ev, turn)

    def _attach_record_identity(self, ev: dict[str, Any], turn: dict[str, Any]) -> dict[str, Any]:
        by_alias = {item.get("alias"): item for item in turn.get("responses") or []}
        slots = {slot.alias: slot for slot in self.slots}
        ev["session_id"] = self.run_id
        ev["turn_id"] = turn.get("index")
        ev["turn_index"] = turn.get("index")
        for model in ev.get("models") or []:
            alias = model.get("alias")
            resp = by_alias.get(alias) or {}
            slot = slots.get(alias)
            model["session_id"] = self.run_id
            model["turn_id"] = turn.get("index")
            model["model_id"] = model.get("model_id") or resp.get("model_id") or (slot.model_id if slot else "")
            model["provider"] = model.get("provider") or resp.get("provider") or (slot.provider if slot else "")
            if "customer_message" not in model:
                model["customer_message"] = turn.get("utterance") or ""
            if model.get("model_response") is None:
                model["model_response"] = resp.get("answer") or resp.get("raw_response")
            model.setdefault("evaluation_status", "OK")
        return ev

    def _error_turn_evaluation(self, turn: dict[str, Any], exc: Exception) -> dict[str, Any]:
        by_alias = {item.get("alias"): item for item in turn.get("responses") or []}
        models = []
        for alias in DISPLAY_ORDER:
            slot = next((item for item in self.slots if item.alias == alias), None)
            resp = by_alias.get(alias) or {}
            models.append(
                {
                    "alias": alias,
                    "provider": (slot.provider if slot else alias.split("_")[0]),
                    "model_id": (slot.model_id if slot else "") or resp.get("model_id") or "",
                    "evaluation_status": EVAL_STATUS_EVALUATION_ERROR,
                    "customer_message": turn.get("utterance") or "",
                    "model_response": resp.get("answer") or resp.get("raw_response"),
                    "api_error": False,
                    "error_type": EVAL_STATUS_EVALUATION_ERROR,
                    "error_message": str(exc),
                    "intent_correct": None,
                    "action_correct": None,
                    "stage_correct": None,
                    "note": "Evaluator exception — not scored as model failure",
                }
            )
        ev = {
            "turn_index": turn.get("index"),
            "turn_id": turn.get("index"),
            "session_id": self.run_id,
            "utterance": turn.get("utterance") or "",
            "kind": turn.get("kind"),
            "ground_truth": "NOT AVAILABLE",
            "models": models,
        }
        return self._attach_record_identity(ev, turn)

    def preview_payloads(self, utterance: str) -> dict[str, str]:
        """Build semantically identical user payloads without calling APIs."""
        customer_history = [t["utterance"] for t in self.turns if t.get("kind") != "opening"]
        extra = {
            "conversation_stage": shared_stage(
                utterance, self.next_turn_index(), self.conversation_mode, customer_history
            )
        }
        return {
            alias: build_user_payload(self.histories[alias], utterance, self.retrieved_context, extra)
            for alias in DISPLAY_ORDER
        }

    def export(self, output_root: Path | None = None) -> Path:
        root = output_root or (self.eval_root / "output" / "interactive_runs" / self.run_id)
        root.mkdir(parents=True, exist_ok=True)
        conversation = {
            "run_id": self.run_id,
            "created_at": self.created_at,
            "mode": (
                "OFFLINE / OUTBOUND INTERACTIVE EVALUATION"
                if self.conversation_mode == "outbound"
                else "OFFLINE / INTERACTIVE EVALUATION"
            ),
            "conversation_mode": self.conversation_mode,
            "opening_mode": self.opening_mode if self.conversation_mode == "outbound" else None,
            "twilio": "DISABLED",
            "retrieved_context": self.retrieved_context,
            "models": [
                {
                    "alias": slot.alias,
                    "provider": slot.provider,
                    "model_id": slot.model_id,
                    "status": slot.status,
                }
                for slot in self.slots
            ],
            "turns": [{"run_id": self.run_id, "turn": t["index"], "timestamp": t["timestamp"], "customer_message": t["utterance"]} for t in self.turns],
        }
        (root / "conversation.json").write_text(json.dumps(conversation, indent=2, ensure_ascii=False), encoding="utf-8")
        with (root / "model_responses.jsonl").open("w", encoding="utf-8") as handle:
            for turn in self.turns:
                for resp in turn["responses"]:
                    row = {"run_id": self.run_id, "turn": turn["index"], "timestamp": turn["timestamp"], **resp}
                    handle.write(json.dumps(row, ensure_ascii=False) + "\n")
        with (root / "evaluations.jsonl").open("w", encoding="utf-8") as handle:
            for row in flatten_evaluation_records(self.evaluations, session_id=self.run_id):
                handle.write(
                    json.dumps({"run_id": self.run_id, **row}, ensure_ascii=False, default=str) + "\n"
                )
        summary = session_summary(self.evaluations)
        csv_path = root / "session_summary.csv"
        fieldnames = ["metric"] + list(DISPLAY_ORDER)
        metrics = [
            ("intent_accuracy", lambda s: s["intent"]["accuracy"]),
            ("intent_precision", lambda s: s["intent"]["precision"]),
            ("intent_recall", lambda s: s["intent"]["recall"]),
            ("intent_macro_f1", lambda s: s["intent"]["f1"]),
            ("action_accuracy", lambda s: s["action"]["accuracy"]),
            ("action_macro_f1", lambda s: s["action"]["f1"]),
            ("factual_accuracy", lambda s: s["factual_accuracy"]),
            ("grounded_pct", lambda s: s["grounded_pct"]),
            ("unsupported_claim_rate", lambda s: s["unsupported_rate"]),
            ("context_accuracy", lambda s: s["context_accuracy"]),
            ("stage_accuracy", lambda s: s["stage_accuracy"]),
            ("visit_sequence_accuracy", lambda s: s.get("visit_sequence_accuracy")),
            ("avg_relevance", lambda s: s["avg_relevance"]),
            ("avg_completeness", lambda s: s["avg_completeness"]),
            ("avg_clarity", lambda s: s["avg_clarity"]),
            ("avg_conversational", lambda s: s["avg_conversational"]),
            ("avg_latency", lambda s: s["avg_latency"]),
            ("p50_latency", lambda s: s["p50_latency"]),
            ("p95_latency", lambda s: s["p95_latency"]),
        ]
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for name, getter in metrics:
                writer.writerow({"metric": name, **{alias: getter(summary[alias]) for alias in DISPLAY_ORDER}})
            opening_rows = _opening_metric_rows(self.evaluations)
            for row in opening_rows:
                writer.writerow(row)
        self._write_excel(root / "session_summary.xlsx", summary)
        self.exported_dir = root
        return root

    def _write_excel(self, path: Path, summary: dict[str, dict[str, Any]]) -> None:
        wb = Workbook()
        conv = wb.active
        conv.title = "Conversation"
        conv.append(["Run ID", "Turn", "Timestamp", "Customer Message"])
        for cell in conv[1]:
            cell.font = Font(bold=True)
        for turn in self.turns:
            conv.append([self.run_id, turn["index"], turn["timestamp"], turn["utterance"]])

        resp_ws = wb.create_sheet("Responses")
        resp_ws.append(
            [
                "Run ID",
                "Turn",
                "Model",
                "Provider",
                "Response",
                "Intent",
                "Action",
                "Stage",
                "Language",
                "Latency",
                "Input Tokens",
                "Output Tokens",
                "Error",
            ]
        )
        for cell in resp_ws[1]:
            cell.font = Font(bold=True)
        for turn in self.turns:
            for resp in turn["responses"]:
                resp_ws.append(
                    [
                        self.run_id,
                        turn["index"],
                        DISPLAY_NAMES.get(resp["alias"], resp["alias"]),
                        resp.get("provider"),
                        resp.get("answer") or resp.get("raw_response"),
                        resp.get("intent"),
                        resp.get("action"),
                        resp.get("stage"),
                        resp.get("language"),
                        resp.get("latency_ms"),
                        resp.get("input_tokens"),
                        resp.get("output_tokens"),
                        _safe_error(resp.get("error_message") or resp.get("error_type")),
                    ]
                )

        ev_ws = wb.create_sheet("Evaluation")
        ev_ws.append(
            [
                "Run ID",
                "Turn",
                "Model",
                "Intent Correct",
                "Precision",
                "Recall",
                "F1",
                "Action Correct",
                "Factual Accuracy",
                "Grounded",
                "Unsupported Claims",
                "Hallucination",
                "Context Correct",
                "Stage Correct",
                "Relevance",
                "Completeness",
                "Clarity",
                "Conversational Quality",
            ]
        )
        for cell in ev_ws[1]:
            cell.font = Font(bold=True)
        for ev in self.evaluations:
            for model in ev.get("models") or []:
                ctx = bool(model.get("context_used")) and not model.get("context_error")
                ev_ws.append(
                    [
                        self.run_id,
                        ev.get("turn_index"),
                        DISPLAY_NAMES.get(model.get("alias"), model.get("alias")),
                        model.get("intent_correct"),
                        None,
                        None,
                        None,
                        model.get("action_correct"),
                        model.get("factual_accuracy"),
                        model.get("grounded"),
                        json.dumps(model.get("unsupported_claims") or [], ensure_ascii=False),
                        model.get("hallucination"),
                        ctx if not model.get("api_error") else None,
                        model.get("stage_correct"),
                        model.get("relevance"),
                        model.get("completeness"),
                        model.get("clarity"),
                        model.get("conversational"),
                    ]
                )

        sum_ws = wb.create_sheet("Session Summary")
        sum_ws.append(["Metric"] + [DISPLAY_NAMES[a] for a in DISPLAY_ORDER])
        for cell in sum_ws[1]:
            cell.font = Font(bold=True)
        rows = [
            ("Intent Accuracy", lambda s: s["intent"]["accuracy"]),
            ("Intent Precision", lambda s: s["intent"]["precision"]),
            ("Intent Recall", lambda s: s["intent"]["recall"]),
            ("Intent Macro F1", lambda s: s["intent"]["f1"]),
            ("Action Accuracy", lambda s: s["action"]["accuracy"]),
            ("Action Macro F1", lambda s: s["action"]["f1"]),
            ("Factual Accuracy", lambda s: s["factual_accuracy"]),
            ("Grounded Response %", lambda s: s["grounded_pct"]),
            ("Unsupported Claim Rate", lambda s: s["unsupported_rate"]),
            ("Context Accuracy", lambda s: s["context_accuracy"]),
            ("Stage Accuracy", lambda s: s["stage_accuracy"]),
            ("Visit Sequence Accuracy", lambda s: s.get("visit_sequence_accuracy")),
            ("Avg Relevance", lambda s: s["avg_relevance"]),
            ("Avg Completeness", lambda s: s["avg_completeness"]),
            ("Avg Clarity", lambda s: s["avg_clarity"]),
            ("Avg Conversational Quality", lambda s: s["avg_conversational"]),
            ("Avg Latency", lambda s: s["avg_latency"]),
            ("P50 Latency", lambda s: s["p50_latency"]),
            ("P95 Latency", lambda s: s["p95_latency"]),
        ]
        for label, getter in rows:
            sum_ws.append([label] + [getter(summary[alias]) for alias in DISPLAY_ORDER])
        wb.save(path)

    def check_models(self) -> dict[str, Any]:
        if self.dry_run:
            return {
                "check_models": True,
                "dry_run": True,
                "results": [
                    {
                        "alias": slot.alias,
                        "provider": slot.provider,
                        "model_id": slot.model_id,
                        "status": "DRY_RUN",
                        "error_message": None,
                    }
                    for slot in self.slots
                ],
            }
        results = []
        for slot in self.slots:
            results.append(self._probe_slot(slot))
        return {"check_models": True, "results": results, "dry_run": False}

    def _probe_slot(self, slot: ModelSlot) -> dict[str, Any]:
        runtime = slot.runtime or {}
        env_name = runtime.get("api_key_env") or "API_KEY"
        if not runtime.get("api_key"):
            return {
                "alias": slot.alias,
                "provider": slot.provider,
                "model_id": slot.model_id,
                "status": "UNAVAILABLE",
                "auth": "NOT CONFIGURED",
                "model_check": "UNAVAILABLE",
                "error_type": "MISSING_API_KEY",
                "error_message": f"{env_name}: NOT CONFIGURED",
            }
        if not runtime.get("model"):
            return {
                "alias": slot.alias,
                "provider": slot.provider,
                "model_id": slot.model_id,
                "status": "UNAVAILABLE",
                "auth": "NOT CONFIGURED",
                "model_check": "UNAVAILABLE",
                "error_type": "MISSING_MODEL_ID",
                "error_message": f"missing model id ({runtime.get('model_env')})",
            }
        timeout = int(runtime.get("timeout") or 30)
        if slot.provider == "claude":
            probe = check_claude_model(
                runtime["api_key"],
                runtime["model"],
                runtime.get("models_url") or "https://api.anthropic.com/v1/models",
                timeout=timeout,
            )
        elif slot.provider == "sarvam":
            probe = check_sarvam_model(runtime["api_key"], runtime["model"], runtime["base_url"], timeout=timeout)
        elif slot.provider == "deepseek":
            probe = check_deepseek_model(
                runtime["api_key"],
                runtime["model"],
                runtime["base_url"],
                timeout=timeout,
                thinking_mode=runtime.get("thinking_mode") or "disabled",
            )
        else:
            probe = {
                "status": "UNAVAILABLE",
                "auth": "FAIL",
                "model": "UNAVAILABLE",
                "error_type": "API_ERROR",
                "error_message": "unsupported provider",
            }
        return {
            "alias": slot.alias,
            "provider": slot.provider,
            "model_id": slot.model_id,
            "status": probe.get("status"),
            "auth": probe.get("auth"),
            "model_check": probe.get("model"),
            "error_type": probe.get("error_type"),
            "error_message": probe.get("error_message"),
        }


def format_check_report(result: dict[str, Any]) -> str:
    lines = ["MODEL CHECK", "=" * 60]
    for row in result.get("results") or []:
        lines.append(f"Provider: {row.get('provider')}")
        lines.append(f"Logical Model Name: {row.get('alias')}")
        lines.append(f"Configured Model ID: {row.get('model_id')}")
        lines.append(f"API availability: {row.get('status')}")
        if row.get("auth"):
            lines.append(f"Authentication: {row.get('auth')}")
        if row.get("model_check"):
            lines.append(f"Model check: {row.get('model_check')}")
        if row.get("error_type") or row.get("error_message"):
            lines.append(f"Error: {redact(row.get('error_message') or row.get('error_type'))}")
        lines.append("-" * 40)
    if result.get("dry_run"):
        lines.append("Dry-run: no external API calls.")
    return "\n".join(lines)
