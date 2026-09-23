"""LLM-EVAL.5.1 — master evaluation report aggregation (no production changes)."""

from __future__ import annotations

import json
from pathlib import Path

from reporting.aggregator import build_report_data, model_eval_observations
from reporting.data_loader import CLASS_FRAMEWORK, load_all, load_interactive_run
from reporting.master_report import generate
from reporting.validation import reconcile


def _write_run(folder: Path, run_id: str, *, fake: bool, language_track: str | None, turns: list[str]) -> None:
    folder.mkdir(parents=True)
    created = "2026-09-23T08:48:57.959346+00:00"
    conv = {
        "run_id": run_id,
        "created_at": created,
        "mode": "OFFLINE / OUTBOUND INTERACTIVE EVALUATION",
        "conversation_mode": "outbound",
        "language_track": language_track,
        "opening_mode": "generated",
        "twilio": "DISABLED",
        "models": [
            {
                "alias": "sarvam_conversational",
                "provider": "sarvam",
                "model_id": "sarvam-105b-conversations",
                "status": "AVAILABLE",
            }
        ],
        "turns": [
            {"run_id": run_id, "turn": i + 1, "timestamp": created, "customer_message": msg}
            for i, msg in enumerate([""] + turns)
        ],
    }
    (folder / "conversation.json").write_text(json.dumps(conv), encoding="utf-8")
    records = []
    for i, msg in enumerate([""] + turns, start=1):
        model_id = "fake-sarvam_conversational" if fake and i > 1 else "sarvam-105b-conversations"
        labelled = bool(msg.strip())
        records.append(
            {
                "run_id": run_id,
                "alias": "sarvam_conversational",
                "provider": "fake" if fake and i > 1 else "sarvam",
                "model_id": model_id,
                "evaluation_status": "OK",
                "customer_message": msg,
                "expected_intent": "QUALIFICATION" if labelled else None,
                "predicted_intent": "QUALIFICATION" if labelled else "GENERAL",
                "intent_correct": True if labelled else None,
                "expected_action": "QUALIFY" if labelled else None,
                "predicted_action": "QUALIFY" if labelled else "ANSWER",
                "action_correct": True if labelled else None,
                "ground_truth_applicable": labelled,
                "factual_accuracy": 1.0,
                "grounded": True,
                "unsupported_claims": [],
                "latency_ms": 20 if fake else 1200,
                "input_tokens": 8 if fake else 900,
                "output_tokens": 4 if fake else 80,
                "total_tokens": 12 if fake else 980,
                "turn_id": i,
                "model_alias": "sarvam_conversational",
            }
        )
    (folder / "evaluations.jsonl").write_text(
        "".join(json.dumps(row) + "\n" for row in records),
        encoding="utf-8",
    )
    (folder / "session_summary.csv").write_text("metric,sarvam_conversational\nintent_accuracy,1.0\n", encoding="utf-8")


def test_fake_provider_run_is_framework_not_model_trial(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    run_dir = root / "output" / "interactive_runs" / "20260923_092857"
    _write_run(run_dir, "20260923_092857", fake=True, language_track="english", turns=["Yes", "My budget is around 2 crore"])
    loaded = load_all(root)
    assert loaded["runs"][0].classification == CLASS_FRAMEWORK
    assert model_eval_observations(loaded["runs"]) == []


def test_session_export_copy_not_double_counted(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    run_dir = root / "output" / "interactive_runs" / "20260923_084857"
    _write_run(run_dir, "20260923_084857", fake=False, language_track=None, turns=["Yes"] * 11)
    (run_dir / "session_summary.xlsx").write_bytes(b"not-a-real-xlsx")
    loaded = load_all(root)
    data = build_report_data(loaded)
    assert len(loaded["runs"]) == 1
    assert data["scale"]["unique_evaluation_runs"] == 1
    assert len(data["observations"]) == 12


def test_batch_and_interactive_are_distinct_runs(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    out = root / "output"
    out.mkdir(parents=True)
    csv_path = out / "evaluation_results.csv"
    csv_path.write_text(
        "run_id,test_id,category,language,provider,model_alias,model_id,customer_utterance,"
        "expected_intent,predicted_intent,intent_pass,expected_action,predicted_action,action_pass,"
        "fact_accuracy,grounded,hallucination,latency_ms,input_tokens,output_tokens,total_tokens,error_type\n"
        "20260918_093501,L001,PROJECT_FAQ,en,sarvam,sarvam_conversational,sarvam-105b-conversations,"
        "What amenities?,AMENITIES_QUERY,AMENITIES_QUERY,True,ANSWER,ANSWER,True,1.0,True,False,1000,500,50,550,\n",
        encoding="utf-8",
    )
    run_dir = out / "interactive_runs" / "20260923_084857"
    _write_run(run_dir, "20260923_084857", fake=False, language_track=None, turns=["Yes"] * 11)
    loaded = load_all(root)
    ids = {r.run_id for r in loaded["runs"]}
    assert ids == {"20260918_093501", "20260923_084857"}
    data = build_report_data(loaded)
    recon = reconcile(data)
    assert recon["checks"]["run_count"] == "PASS"
    assert recon["checks"]["model_turn_count"] == "PASS"
    assert recon["checks"]["no_duplicate_run_id"] == "PASS"


def test_hindi_exploratory_not_marked_h01_h08(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    run_dir = root / "output" / "interactive_runs" / "20260923_103454"
    _write_run(
        run_dir,
        "20260923_103454",
        fake=False,
        language_track="hindi",
        turns=["Han", "teen BHK", "Visit kaise karna na?"],
    )
    bundle, _ = load_interactive_run(run_dir, root)
    assert bundle.classification != CLASS_FRAMEWORK
    assert "H01" not in bundle.status
    assert bundle.evidence_class == "EXPLORATORY INTERACTIVE"


def test_eval1_baseline_uses_folder_run_id(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    run_dir = root / "output" / "runs" / "20260918_071117"
    run_dir.mkdir(parents=True)
    (run_dir / "evaluation_results.csv").write_text(
        "test_id,category,language,provider,model,customer_utterance,expected_intent,predicted_intent,"
        "intent_pass,expected_action,predicted_action,action_pass,fact_accuracy,grounded,latency_ms,"
        "input_tokens,output_tokens,total_tokens\n"
        "L001,PROJECT_FAQ,en,sarvam,sarvam-105b-conversations,What amenities?,AMENITIES_QUERY,"
        "AMENITIES_QUERY,True,ANSWER,ANSWER,True,1.0,True,1000,500,50,550\n"
        "L001,PROJECT_FAQ,en,claude,claude-sonnet-4-20250514,What amenities?,AMENITIES_QUERY,"
        "AMENITIES_QUERY,True,ANSWER,ANSWER,True,1.0,True,2000,600,80,680\n",
        encoding="utf-8",
    )
    loaded = load_all(root)
    assert {r.run_id for r in loaded["runs"]} == {"20260918_071117"}
    assert "LLM-EVAL.1" in loaded["runs"][0].status
    data = build_report_data(loaded)
    names = {m["display_model"] for m in data["model_metrics"]}
    assert "Sarvam 105B Conversations" in names
    assert "claude-sonnet-4-20250514" in names


def test_generate_workbook_reconciliation(tmp_path: Path) -> None:
    root = tmp_path / "LLM_Evaluation"
    (root / "tests").mkdir(parents=True)
    (root / "tests" / "test_eval_framework.py").write_text(
        "def test_dataset_loading():\n    assert True\n",
        encoding="utf-8",
    )
    run_dir = root / "output" / "interactive_runs" / "20260923_084857"
    _write_run(run_dir, "20260923_084857", fake=False, language_track=None, turns=["Yes"] * 11)
    result = generate(root, root / "output" / "AI_Calling_Agent_Master_Evaluation_Report.xlsx")
    assert result["workbook"].exists()
    assert result["reconciliation"]["checks"]["run_count"] == "PASS"
    assert result["reconciliation"]["checks"]["model_turn_count"] == "PASS"
    assert result["reconciliation"]["checks"]["token_reconciliation"] == "PASS"
    assert result["reconciliation"]["checks"]["cost_reconciliation"] == "PASS"
    assert result["scale"]["framework_tests"] >= 1
    assert result["scale"]["unique_evaluation_runs"] == 1
    from openpyxl import load_workbook

    wb = load_workbook(result["workbook"])
    assert "Executive Summary" in wb.sheetnames
    assert "Evaluation Run Inventory" in wb.sheetnames
    assert "Detailed Model-Turn Results" in wb.sheetnames
    assert "Framework Validation" in wb.sheetnames
    assert "Data Lineage Sources" in wb.sheetnames
