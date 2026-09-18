from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd

SUMMARY_COLUMNS = [
    "display_name",
    "provider",
    "Model Role",
    "Model ID",
    "Availability",
    "Benchmark Cases",
    "Successful Runs",
    "API Errors",
    "Completion Rate",
    "Intent Accuracy",
    "Action Accuracy",
    "Fact Accuracy",
    "Grounded %",
    "Hallucination %",
    "Context Handling %",
    "Language Handling %",
    "Schema Valid %",
    "Stage Appropriateness",
    "Avg Response Quality",
    "Avg Latency",
    "P50 Latency",
    "P95 Latency",
    "Avg Response Words",
    "Input Tokens",
    "Output Tokens",
    "Total Tokens",
    "Estimated Cost",
    "Notes",
]

LEGACY_SUMMARY_COLUMNS = [
    "provider",
    "Test Cases",
    "Intent Accuracy",
    "Action Accuracy",
    "Fact Accuracy",
    "Grounded %",
    "Hallucination %",
    "Context Handling %",
    "Language Handling %",
    "Schema Valid %",
    "Avg Response Quality",
    "Avg Latency",
    "P50 Latency",
    "P95 Latency",
    "Total Tokens",
    "Errors",
]


def _excel_cell(value: Any) -> Any:
    if isinstance(value, (list, dict, tuple, set)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_frame(ws, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in dataframe_to_rows(frame, index=False, header=False):
        ws.append([_excel_cell(item) for item in row])


def write_workbook(
    path: Path,
    summary_rows: list[dict[str, Any]],
    detail_rows: list[dict[str, Any]],
    error_rows: list[dict[str, Any]],
    dataset_rows: list[dict[str, Any]],
    extra_meta: dict[str, Any] | None = None,
    comparison_rows: list[dict[str, Any]] | None = None,
    matrix_rows: list[dict[str, Any]] | None = None,
    latency_rows: list[dict[str, Any]] | None = None,
    config_rows: list[dict[str, Any]] | None = None,
) -> None:
    wb = Workbook()
    summary_df = pd.DataFrame(summary_rows)
    use_v11 = bool(summary_rows) and "display_name" in summary_df.columns

    ws1 = wb.active
    ws1.title = "Executive Summary"
    if extra_meta:
        ws1.append(["run_id", extra_meta.get("run_id")])
        ws1.append(["evaluation_version", extra_meta.get("eval_version") or extra_meta.get("evaluation_version")])
        ws1.append(["dataset_version", extra_meta.get("dataset_version")])
        ws1.append(["prompt_version", extra_meta.get("prompt_version")])
        ws1.append(["baseline", extra_meta.get("baseline", "LLM-EVAL.1.1")])
        ws1.append([])
    cols = SUMMARY_COLUMNS if use_v11 else [c for c in LEGACY_SUMMARY_COLUMNS if c in summary_df.columns or True]
    if use_v11:
        for col in SUMMARY_COLUMNS:
            if col not in summary_df.columns:
                summary_df[col] = None
        summary_df = summary_df[SUMMARY_COLUMNS]
    else:
        for col in LEGACY_SUMMARY_COLUMNS:
            if col not in summary_df.columns:
                summary_df[col] = None
        summary_df = summary_df[LEGACY_SUMMARY_COLUMNS]
    _write_frame(ws1, summary_df)

    ws2 = wb.create_sheet("Provider Comparisons")
    cmp_df = pd.DataFrame(comparison_rows or [])
    _write_frame(ws2, cmp_df if not cmp_df.empty else pd.DataFrame([{"note": "no_comparison"}]))

    ws3 = wb.create_sheet("Detailed Results")
    detail_df = pd.DataFrame(detail_rows)
    _write_frame(ws3, detail_df if not detail_df.empty else pd.DataFrame([{"note": "no_rows"}]))

    ws4 = wb.create_sheet("Model x Test Matrix")
    mat_df = pd.DataFrame(matrix_rows or [])
    _write_frame(ws4, mat_df if not mat_df.empty else pd.DataFrame([{"note": "no_matrix"}]))

    ws5 = wb.create_sheet("Error Analysis")
    err_df = pd.DataFrame(error_rows)
    _write_frame(ws5, err_df if not err_df.empty else pd.DataFrame([{"note": "no_errors"}]))

    ws6 = wb.create_sheet("Multilingual")
    ws6.append(["NOTE", "INITIAL MULTILINGUAL SMOKE TEST — not a statistically sufficient multilingual benchmark"])
    multi = [
        row
        for row in detail_rows
        if str(row.get("language", "")).lower() in {"mixed", "hi", "kn"} or row.get("category") == "MULTILINGUAL"
    ]
    multi_cols = [
        "test_id",
        "language",
        "provider",
        "model_alias",
        "model_id",
        "predicted_language",
        "language_pass",
        "intent_pass",
        "fact_accuracy",
        "grounded",
        "overall_response_quality",
        "answer",
        "notes",
    ]
    multi_df = pd.DataFrame(multi)
    if not multi_df.empty:
        for col in multi_cols:
            if col not in multi_df.columns:
                multi_df[col] = None
        multi_df = multi_df[[c for c in multi_cols if c in multi_df.columns]]
    _write_frame(ws6, multi_df if not multi_df.empty else pd.DataFrame([{"note": "no_multilingual_rows"}]))

    ws7 = wb.create_sheet("Latency & Usage")
    lat_source = latency_rows if latency_rows is not None else summary_rows
    lat_df = pd.DataFrame(lat_source)
    lat_keep = [
        c
        for c in [
            "display_name",
            "Avg Latency",
            "P50 Latency",
            "P95 Latency",
            "Min Latency",
            "Max Latency",
            "Avg Input Tokens",
            "Avg Output Tokens",
            "Total Tokens",
            "Estimated Cost",
        ]
        if c in lat_df.columns
    ]
    _write_frame(ws7, lat_df[lat_keep] if lat_keep and not lat_df.empty else pd.DataFrame([{"note": "no_latency_rows"}]))

    ws8 = wb.create_sheet("Model Configuration")
    cfg_df = pd.DataFrame(config_rows or [])
    _write_frame(ws8, cfg_df if not cfg_df.empty else pd.DataFrame([{"note": "no_config_rows"}]))

    ws9 = wb.create_sheet("Test Dataset")
    ds_df = pd.DataFrame(dataset_rows)
    if "conversation_history" in ds_df.columns:
        ds_df = ds_df.copy()
        ds_df["conversation_history"] = ds_df["conversation_history"].apply(lambda v: str(v))
        ds_df["expected_facts"] = ds_df["expected_facts"].apply(lambda v: str(v))
    _write_frame(ws9, ds_df)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
